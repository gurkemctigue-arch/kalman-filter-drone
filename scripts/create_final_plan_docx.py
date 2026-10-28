"""Create the polished Chinese final plan from docs/FINAL_PLAN.md.

This fallback uses python-docx because the local Node environment does not
provide the docx-js package. It keeps the source of truth in Markdown while
producing a readable Word document for review.
"""

from __future__ import annotations

import re
import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "FINAL_PLAN.md"
OUTPUT = ROOT / "docs" / "最终版毕业设计计划书.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_east_asia_font(style_or_run, font_name: str) -> None:
    """Set the font for both Latin and East Asian text in Word XML.

    python-docx's ``font.name`` only sets the ASCII font in some versions.
    Explicitly setting all four font slots keeps Chinese text consistent when
    the document is opened on a machine with a different default font.
    """

    style_or_run.font.name = font_name
    rpr = style_or_run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), font_name)


def add_code_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["No Spacing"]
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    return paragraph


def new_numbering_instance(document: Document) -> int:
    """Create a fresh decimal-list instance that starts at 1.

    The built-in ``List Number`` style points to one shared numbering
    instance.  Depending on the Word-compatible renderer, separate Markdown
    lists can therefore continue numbering across sections.  A dedicated
    ``w:num`` with an explicit level-0 start override makes each list restart
    deterministically while retaining a real Word numbering definition.
    """

    numbering = document.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        p_style = abstract.find(f".//{qn('w:pStyle')}")
        if p_style is not None and p_style.get(qn("w:val")) == "ListNumber":
            abstract_id = abstract.get(qn("w:abstractNumId"))
            break
    if abstract_id is None:
        raise RuntimeError("The python-docx template has no List Number definition")

    existing_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    num_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def set_numbering(paragraph, num_id: int) -> None:
    """Attach a level-0 numbering instance to a paragraph."""

    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))
    ppr.append(num_pr)


def build(output: Path = OUTPUT) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # python-docx leaves ``w:zoom`` without its required percentage.  Word
    # opens such files, but strict OOXML validators correctly flag it.
    settings = document.settings._element
    zoom = settings.find(qn("w:zoom"))
    if zoom is None:
        zoom = OxmlElement("w:zoom")
        settings.append(zoom)
    zoom.set(qn("w:percent"), "100")

    normal = document.styles["Normal"]
    set_east_asia_font(normal, "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        set_east_asia_font(style, "Noto Sans CJK SC")
        style.font.color.rgb = None
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(13)
    document.styles["Heading 3"].font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(5)

    in_code = False
    code_paragraphs = []
    table = None
    previous_numbered = False
    current_num_id = None
    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            if in_code:
                code_paragraphs = []
            else:
                for paragraph in code_paragraphs[:-1]:
                    paragraph.paragraph_format.keep_with_next = True
            previous_numbered = False
            continue
        if in_code:
            code_paragraphs.append(add_code_paragraph(document, line))
            previous_numbered = False
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        line = line.replace("`", "")
        if numbered:
            if not previous_numbered:
                current_num_id = new_numbering_instance(document)
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(numbered.group(2).replace("`", ""))
            set_numbering(paragraph, current_num_id)
            previous_numbered = True
            continue
        previous_numbered = False

        if line.startswith("# "):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:])
            run.bold = True
            set_east_asia_font(run, "Noto Sans CJK SC")
            run.font.size = Pt(18)
            continue
        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=2)
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:])
            continue
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(set(cell) <= {"-", ":", " "} for cell in cells):
                continue
            if table is None:
                table = document.add_table(rows=1, cols=len(cells))
                table.style = "Table Grid"
                for index, cell in enumerate(cells):
                    table.rows[0].cells[index].text = cell
                    set_cell_shading(table.rows[0].cells[index], "D9EAF7")
            else:
                row = table.add_row().cells
                for index, cell in enumerate(cells):
                    row[index].text = cell
            continue
        if not line:
            table = None
            continue
        document.add_paragraph(line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("RGB-D 无人机目标跟踪与自主跟随系统  ·  ")
    footer_run.font.size = Pt(8)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    build(parser.parse_args().output)
